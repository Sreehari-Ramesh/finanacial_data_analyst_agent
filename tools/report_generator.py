"""Generate Excel and CSV reports from query results."""
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import REPORTS_DIR


def _sanitize_filename(text: str) -> str:
    """Convert a question string to a safe filename fragment."""
    clean = re.sub(r"[^\w\s-]", "", text.lower())
    clean = re.sub(r"[\s_-]+", "_", clean).strip("_")
    return clean[:50]


def _col_width(df: pd.DataFrame, col: str) -> int:
    """Compute a reasonable column width."""
    max_data = df[col].astype(str).str.len().max() if len(df) else 0
    return min(max(len(col), max_data) + 4, 50)


def generate_excel_report(
    rows: list[dict],
    columns: list[str],
    question: str,
    title: str = "",
    sql_query: str = "",
) -> str:
    """Create a formatted Excel report and return the file path."""

    df = pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)
    title = title or f"Report: {question[:60]}"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    slug      = _sanitize_filename(question)
    filename  = REPORTS_DIR / f"report_{timestamp}_{slug}.xlsx"

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        # ── Data sheet ──────────────────────────────────────────────────────
        df.to_excel(writer, sheet_name="Data", startrow=3, index=False)
        ws = writer.sheets["Data"]

        # Title
        ws.cell(1, 1, title).font      = Font(bold=True, size=14)
        ws.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = Font(size=10, italic=True)

        # Header row (row 4 in Excel = startrow=3 + 1-based)
        header_fill   = PatternFill("solid", fgColor="1F4E79")
        header_font   = Font(bold=True, color="FFFFFF")
        header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border   = Border(
            bottom=Side(style="thin", color="FFFFFF"),
            right =Side(style="thin", color="FFFFFF"),
        )

        num_cols = len(df.columns)
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(4, col_idx)
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = header_align
            cell.border    = thin_border

        # Alternate row shading
        alt_fill = PatternFill("solid", fgColor="D9E1F2")
        for row_idx in range(5, 5 + len(df)):
            if (row_idx % 2) == 0:
                for col_idx in range(1, num_cols + 1):
                    ws.cell(row_idx, col_idx).fill = alt_fill

        # Column widths
        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = _col_width(df, col_name)

        ws.freeze_panes = "A5"

        # ── Summary sheet ────────────────────────────────────────────────────
        ws2 = writer.book.create_sheet("Summary")
        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 30

        hdr = Font(bold=True, size=13, color="1F4E79")
        ws2.cell(1, 1, "Report Summary").font = hdr
        ws2.cell(2, 1, "Question").font       = Font(bold=True)
        ws2.cell(2, 2, question)
        ws2.cell(3, 1, "Generated").font      = Font(bold=True)
        ws2.cell(3, 2, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ws2.cell(4, 1, "Total Rows").font     = Font(bold=True)
        ws2.cell(4, 2, len(df))

        if not df.empty:
            ws2.cell(6, 1, "Column Statistics").font = Font(bold=True, size=12)
            ws2.cell(7, 1, "Column").font  = Font(bold=True)
            ws2.cell(7, 2, "Type").font    = Font(bold=True)
            ws2.cell(7, 3, "Non-Null").font = Font(bold=True)
            ws2.cell(7, 4, "Unique").font  = Font(bold=True)
            for col_idx in range(1, 5):
                ws2.cell(7, col_idx).fill = PatternFill("solid", fgColor="1F4E79")
                ws2.cell(7, col_idx).font = Font(bold=True, color="FFFFFF")

            for r, col in enumerate(df.columns, start=8):
                ws2.cell(r, 1, col)
                ws2.cell(r, 2, str(df[col].dtype))
                ws2.cell(r, 3, int(df[col].notna().sum()))
                ws2.cell(r, 4, int(df[col].nunique()))

        # ── SQL sheet ────────────────────────────────────────────────────────
        if sql_query:
            ws3 = writer.book.create_sheet("SQL Query")
            ws3.column_dimensions["A"].width = 100
            ws3.cell(1, 1, "SQL Query Used").font = Font(bold=True, size=12, color="1F4E79")
            for i, line in enumerate(sql_query.splitlines(), start=2):
                ws3.cell(i, 1, line).font = Font(name="Courier New", size=10)

    return str(filename)


def generate_csv_report(
    rows: list[dict],
    columns: list[str],
    question: str,
) -> str:
    """Create a CSV file and return the file path."""
    df = pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    slug      = _sanitize_filename(question)
    filename  = REPORTS_DIR / f"report_{timestamp}_{slug}.csv"

    df.to_csv(filename, index=False)
    return str(filename)
